#excel_exporter.py

"""Экспорт подходящих вакансий в Excel."""

from __future__ import annotations

import logging
import threading
from datetime import datetime
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

logger = logging.getLogger("resume_matcher")

PROJECT_ROOT = Path(__file__).resolve().parent
DEFAULT_XLSX_PATH = PROJECT_ROOT / "suitable_vacancies.xlsx"

HEADERS: tuple[str, ...] = (
    "№",
    "Дата парсинга",
    "Источник",
    "Название вакансии",
    "Ссылка",
    "Совпадение (%)",
    "Плюсы",
    "Минусы",
    "Выжимка",
    "Сопроводительное письмо",
    "Статус отклика",
)

# Ширина колонок в символах Excel (колонка с письмом — самая широкая)
COLUMN_WIDTHS: dict[str, float] = {
    "A": 6,
    "B": 18,
    "C": 24,
    "D": 42,
    "E": 38,
    "F": 16,
    "G": 42,
    "H": 42,
    "I": 48,
    "J": 70,
    "K": 20,
}

DEFAULT_STATUS = "Новый"

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
_WRAP_ALIGN = Alignment(vertical="top", wrap_text=True, horizontal="left")
_CENTER_ALIGN = Alignment(vertical="top", horizontal="center", wrap_text=True)
_THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)


class ExcelExporter:
    """Создаёт suitable_vacancies.xlsx и дописывает в него новые вакансии."""

    def __init__(self, path: Path | str | None = None) -> None:
        self.path = Path(path) if path else DEFAULT_XLSX_PATH
        self._lock = threading.Lock()

    def add_vacancy(
        self,
        *,
        source: str,
        title: str,
        url: str = "",
        match_percentage: int,
        pros: Sequence[str] | str = (),
        cons: Sequence[str] | str = (),
        summary: str = "",
        cover_letter: str = "",
        status: str = DEFAULT_STATUS,
        parsed_at: datetime | None = None,
    ) -> int:
        """
        Добавляет строку в конец таблицы.

        Возвращает порядковый номер записи («№»).
        """
        parsed_at = parsed_at or datetime.now()
        with self._lock:
            workbook = self._open_or_create()
            sheet = workbook.active
            self._ensure_layout(sheet)

            row_number = self._next_row(sheet)
            serial = self._next_serial(sheet, row_number)

            values = (
                serial,
                parsed_at.strftime("%Y-%m-%d %H:%M"),
                source.strip(),
                title.strip(),
                url.strip(),
                int(match_percentage),
                self._join_list(pros),
                self._join_list(cons),
                summary.strip(),
                cover_letter.strip(),
                status.strip() or DEFAULT_STATUS,
            )
            for col, value in enumerate(values, start=1):
                cell = sheet.cell(row=row_number, column=col, value=value)
                self._style_data_cell(cell, col)

            if url.strip():
                link_cell = sheet.cell(row=row_number, column=5)
                link_cell.hyperlink = url.strip()
                link_cell.font = Font(color="0563C1", underline="single")

            # Высокая строка, чтобы письмо и списки плюсов/минусов читались
            sheet.row_dimensions[row_number].height = 90
            self._refresh_filter(sheet, row_number)

            workbook.save(self.path)
            workbook.close()
            logger.info(
                "Вакансия записана в %s: №%s «%s» (%s%%)",
                self.path.name,
                serial,
                title.strip(),
                match_percentage,
            )
            return serial

    def _open_or_create(self) -> Workbook:
        if self.path.exists():
            try:
                workbook = load_workbook(self.path)
            except Exception as exc:
                logger.error("Не удалось открыть %s: %s", self.path, exc)
                raise
            if workbook.active is None:
                workbook.create_sheet("Вакансии")
            return workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Вакансии"
        self._write_headers(sheet)
        logger.info("Создан файл %s", self.path)
        return workbook

    def _ensure_layout(self, sheet: Worksheet) -> None:
        """Если файл пустой или без шапки — создаём заголовки заново."""
        first_row = [sheet.cell(1, col).value for col in range(1, len(HEADERS) + 1)]
        if first_row != list(HEADERS):
            if any(first_row):
                logger.warning(
                    "Заголовки в %s не совпадают с ожидаемыми — оставляю как есть, дописываю строки.",
                    self.path.name,
                )
            else:
                self._write_headers(sheet)
        self._apply_column_widths(sheet)
        sheet.freeze_panes = "A2"
        sheet.row_dimensions[1].height = 24

    def _write_headers(self, sheet: Worksheet) -> None:
        for col, header in enumerate(HEADERS, start=1):
            cell = sheet.cell(row=1, column=col, value=header)
            cell.font = _HEADER_FONT
            cell.fill = _HEADER_FILL
            cell.alignment = _HEADER_ALIGN
            cell.border = _THIN_BORDER
        self._apply_column_widths(sheet)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}1"

    def _apply_column_widths(self, sheet: Worksheet) -> None:
        for letter, width in COLUMN_WIDTHS.items():
            sheet.column_dimensions[letter].width = width

    def _next_row(self, sheet: Worksheet) -> int:
        return sheet.max_row + 1 if sheet.max_row >= 1 else 2

    def _next_serial(self, sheet: Worksheet, new_row: int) -> int:
        """Берём максимальный «№» из колонки A, чтобы нумерация не сбивалась."""
        last_serial = 0
        for row in range(2, new_row):
            value = sheet.cell(row=row, column=1).value
            if isinstance(value, int):
                last_serial = max(last_serial, value)
            elif isinstance(value, str) and value.strip().isdigit():
                last_serial = max(last_serial, int(value.strip()))
        return last_serial + 1

    def _refresh_filter(self, sheet: Worksheet, last_row: int) -> None:
        last_col = get_column_letter(len(HEADERS))
        sheet.auto_filter.ref = f"A1:{last_col}{max(last_row, 1)}"

    @staticmethod
    def _style_data_cell(cell, column: int) -> None:
        cell.border = _THIN_BORDER
        if column in (1, 6, 11):
            cell.alignment = _CENTER_ALIGN
        else:
            cell.alignment = _WRAP_ALIGN

    @staticmethod
    def _join_list(items: Sequence[str] | str | Iterable[str]) -> str:
        if isinstance(items, str):
            return items.strip()
        lines = [item.strip() for item in items if item and str(item).strip()]
        if not lines:
            return ""
        return "\n".join(f"• {line}" for line in lines)
